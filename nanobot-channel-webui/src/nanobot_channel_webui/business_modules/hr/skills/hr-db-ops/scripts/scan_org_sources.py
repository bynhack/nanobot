#!/usr/bin/env python3
import argparse
import json
import re
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook
try:
    import xlrd
except ImportError:
    xlrd = None


SKIP_SHEETS = {"WpsReserved_CellImgList"}
PACKAGE_DIRS = {"人事相关资料（传媒、红人、短视频、酒肆）", "人事系统资料-郭娟"}


def main():
    parser = argparse.ArgumentParser(description="Scan HR Excel files for company and department seeds.")
    parser.add_argument("--root", default=".", help="Workspace or import folder to scan. Defaults to the current working directory.")
    parser.add_argument("--json", action="store_true", help="Print full JSON instead of a business summary.")
    args = parser.parse_args()

    root = Path(args.root)
    result = scan_org_sources(root)
    if args.json:
        print(json.dumps(result, ensure_ascii=False, indent=2))
    else:
        print_summary(result)


def scan_org_sources(root):
    files = sorted(
        p for p in root.rglob("*")
        if p.is_file()
        and p.suffix.lower() in {".xlsx", ".xls"}
        and ".nanobot" not in p.parts
        and "skills" not in p.parts
        and "node_modules" not in p.parts
    )

    companies = {}
    unreadable = []
    scanned_files = []

    for path in files:
        file_kind = classify_file(path)
        if file_kind not in {"roster_active", "roster_resigned", "personnel_change"}:
            continue

        company_key = company_from_path(root, path)
        if not company_key:
            continue

        company = companies.setdefault(company_key, {
            "name": company_key,
            "source_files": set(),
            "departments": defaultdict(lambda: {"sources": set(), "examples": set()}),
            "observed_company_labels": defaultdict(int),
            "notes": set(),
        })
        company["source_files"].add(str(path.relative_to(root)))

        try:
            rows = scan_excel(path, file_kind)
        except Exception as exc:
            unreadable.append({
                "file": str(path.relative_to(root)),
                "reason": repr(exc),
                "kind": file_kind,
                "company": company_key,
            })
            company["notes"].add("存在读取失败文件，部门可能不完整")
            continue

        scanned_files.append({
            "file": str(path.relative_to(root)),
            "kind": file_kind,
            "company": company_key,
            "rows": len(rows),
        })

        for row in rows:
            dept = clean_text(row.get("department"))
            person = clean_text(row.get("name"))
            observed_company = clean_text(row.get("company"))

            if observed_company and not is_noise(observed_company):
                company["observed_company_labels"][observed_company] += 1
            # Department dictionary must come from rosters only. Personnel
            # change sheets are useful for later employee-change records, but
            # their department cells can contain jobs, section labels, or
            # temporary wording that should not seed departments.
            if file_kind in {"roster_active", "roster_resigned"} and dept and not is_noise(dept):
                entry = company["departments"][dept]
                entry["sources"].add(str(path.relative_to(root)))
                if person:
                    entry["examples"].add(person)

    normalized_companies = []
    for company in companies.values():
        normalized_companies.append({
            "name": company["name"],
            "source_files": sorted(company["source_files"]),
            "departments": [
                {
                    "name": name,
                    "source_count": len(info["sources"]),
                    "example_employees": sorted(info["examples"])[:5],
                }
                for name, info in sorted(company["departments"].items())
            ],
            "observed_company_labels": dict(sorted(company["observed_company_labels"].items())),
            "notes": sorted(company["notes"]),
        })

    return {
        "root": str(root),
        "company_count": len(normalized_companies),
        "department_count": sum(len(c["departments"]) for c in normalized_companies),
        "companies": sorted(normalized_companies, key=lambda item: item["name"]),
        "scanned_files": scanned_files,
        "unreadable_files": unreadable,
    }


def scan_excel(path, file_kind):
    if path.suffix.lower() == ".xls":
        return scan_xls(path, file_kind)
    return scan_xlsx(path, file_kind)


def scan_xlsx(path, file_kind):
    workbook = load_workbook(path, read_only=True, data_only=True)
    rows = []
    try:
        for sheet in workbook.worksheets:
            if sheet.title in SKIP_SHEETS:
                continue
            if sheet.sheet_state != "visible":
                continue
            if file_kind in {"roster_active", "roster_resigned"}:
                rows.extend(scan_roster_sheet(sheet))
            elif file_kind == "personnel_change":
                rows.extend(scan_personnel_change_sheet(sheet))
    finally:
        workbook.close()
    return rows


def scan_xls(path, file_kind):
    if xlrd is None:
        raise RuntimeError("xlrd is required to parse legacy .xls files")

    workbook = xlrd.open_workbook(str(path), formatting_info=True)
    rows = []
    for sheet in workbook.sheets():
        if getattr(sheet, "visibility", 0) != 0:
            continue
        if sheet.name in SKIP_SHEETS:
            continue
        adapter = XlsSheetAdapter(sheet)
        if file_kind in {"roster_active", "roster_resigned"}:
            rows.extend(scan_roster_sheet(adapter))
        elif file_kind == "personnel_change":
            rows.extend(scan_personnel_change_sheet(adapter))
    return rows


class XlsSheetAdapter:
    def __init__(self, sheet):
        self.title = sheet.name
        self.sheet = sheet

    def iter_rows(self, min_row=1, values_only=True):
        del values_only
        start = max(min_row - 1, 0)
        for idx in range(start, self.sheet.nrows):
            yield self.sheet.row_values(idx)


def scan_roster_sheet(sheet):
    header_index = None
    header = None
    for idx, values in enumerate(iter_values(sheet, max_cols=60), start=1):
        normalized = [clean_header(v) for v in values]
        if "姓名" in normalized and any(h in normalized for h in ["部门", "部门名称"]):
            header_index = idx
            header = normalized
            break
        if idx > 12:
            break
    if not header:
        return []

    dept_cols = department_indexes(header)
    name_col = first_index(header, ["姓名"])
    company_col = first_index(header, ["公司", "所属公司", "合同公司"])
    rows = []
    current_department_levels = []
    for values in iter_values(sheet, min_row=header_index + 1, max_cols=60):
        name = value_at(values, name_col)
        raw_levels = [clean_text(value_at(values, col)) for col in dept_cols]
        if any(raw_levels):
            current_department_levels = merge_department_levels(current_department_levels, raw_levels)
        dept = build_department_name(current_department_levels)
        if not clean_text(name) and not clean_text(dept):
            continue
        if is_section_row(name) or is_section_row(dept):
            continue
        rows.append({
            "company": value_at(values, company_col),
            "department": dept,
            "name": name,
        })
    return rows


def scan_personnel_change_sheet(sheet):
    rows = []
    for idx, values in enumerate(iter_values(sheet, max_cols=30), start=1):
        normalized = [clean_header(v) for v in values]
        if "公司" not in normalized or "姓名" not in normalized or "部门" not in normalized:
            continue

        company_col = first_index(normalized, ["公司"])
        name_col = first_index(normalized, ["姓名"])
        dept_col = first_index(normalized, ["部门"])
        for data in iter_values(sheet, min_row=idx + 1, max_cols=30):
            name = clean_text(value_at(data, name_col))
            dept = clean_text(value_at(data, dept_col))
            company = clean_text(value_at(data, company_col))
            if not name and not dept and not company:
                continue
            if is_section_row(company) or is_section_row(name) or is_noise(company):
                continue
            rows.append({
                "company": company,
                "department": dept,
                "name": name,
            })
        break
    return rows


def iter_values(sheet, min_row=1, max_cols=60):
    for row in sheet.iter_rows(min_row=min_row, values_only=True):
        yield list(row[:max_cols])


def classify_file(path):
    name = path.name
    parent = path.parent.name
    if "花名册" in name and any(token in name for token in ["在职", "员工在职", "人员花名册", "在职员工", "在职人员"]):
        return "roster_active"
    if "离职" in name and "花名册" in name:
        return "roster_resigned"
    if "离职员工" in name or "离职人员" in name:
        return "roster_resigned"
    if "异动" in name or parent == "人员异动表":
        return "personnel_change"
    return "other"


def company_from_path(root, path):
    rel_parts = path.relative_to(root).parts
    for idx, part in enumerate(rel_parts):
        if part in PACKAGE_DIRS and idx + 1 < len(rel_parts):
            return part_to_company_name(rel_parts[idx + 1])
    return None


def part_to_company_name(part):
    # Keep project suffixes in the preview for human confirmation. They may map
    # to either true companies or internal project departments.
    return clean_text(part)


def clean_header(value):
    return re.sub(r"\s+", "", clean_text(value))


def clean_text(value):
    if value is None:
        return ""
    text = str(value).strip()
    if text in {"", "/", "\\", "nan", "None"}:
        return ""
    return re.sub(r"\s+", " ", text)


def first_index(values, candidates):
    for candidate in candidates:
        if candidate in values:
            return values.index(candidate)
    return None


def department_indexes(header):
    indexes = [idx for idx, value in enumerate(header) if value in {"部门", "部门名称"}]
    if indexes:
        first = indexes[0]
        # Some rosters use merged multi-level department headers. In the read
        # values this appears as "部门", "", "" before position/name.
        extra = []
        for idx in range(first + 1, min(first + 3, len(header))):
            if header[idx] == "":
                extra.append(idx)
            else:
                break
        return [first, *extra]
    return []


def merge_department_levels(current, incoming):
    levels = list(current) if current else ["", "", ""]
    if len(levels) < len(incoming):
        levels.extend([""] * (len(incoming) - len(levels)))

    for idx, value in enumerate(incoming):
        if value:
            levels[idx] = value
            for child_idx in range(idx + 1, len(levels)):
                if child_idx < len(incoming) and incoming[child_idx]:
                    levels[child_idx] = incoming[child_idx]
                elif idx in {0, 1}:
                    levels[child_idx] = ""
    return levels


def build_department_name(levels):
    cleaned = [level for level in levels if level and not is_noise(level)]
    if not cleaned:
        return ""
    if cleaned[0] == "娱乐运营部" and len(cleaned) >= 2:
        return "-".join(cleaned[1:])
    return cleaned[-1]


def value_at(values, index):
    if index is None or index >= len(values):
        return ""
    return values[index]


def is_noise(value):
    text = clean_text(value)
    return text in {
        "部门",
        "入职人员",
        "离职人员",
        "转正人员",
        "调薪人员",
        "调岗人员",
        "晋升人员",
        "岗位异动人员",
        "薪资晋档人员",
        "社保",
        "无",
        "暂无",
    }


def is_section_row(value):
    text = clean_text(value)
    return text in {
        "一",
        "二",
        "三",
        "四",
        "五",
        "入职人员",
        "离职人员",
        "转正人员",
        "调薪人员",
        "调岗人员",
        "晋升人员",
        "岗位异动人员",
        "薪资晋档人员",
        "社保",
    }


def print_summary(result):
    print(f"公司/项目主体: {result['company_count']}")
    print(f"部门候选: {result['department_count']}")
    print(f"已扫描基础文件: {len(result['scanned_files'])}")
    print(f"未解析基础文件: {len(result['unreadable_files'])}")
    print()
    for company in result["companies"]:
        departments = [item["name"] for item in company["departments"]]
        print(company["name"])
        print(f"  部门({len(departments)}): {', '.join(departments) if departments else '无'}")
        labels = company["observed_company_labels"]
        if labels:
            label_text = ", ".join(f"{k}({v})" for k, v in labels.items())
            print(f"  表内公司写法: {label_text}")
        if company["notes"]:
            print(f"  注意: {'; '.join(company['notes'])}")
        print()

    if result["unreadable_files"]:
        print("未解析文件:")
        for item in result["unreadable_files"]:
            print(f"- {item['file']}：{item['reason']}")


if __name__ == "__main__":
    main()
